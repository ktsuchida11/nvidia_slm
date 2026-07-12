"""PPTX/PDF/DOCX → S1入力(jsonl) の軽量コンバータ（CPU・テキスト/表/ノートのみ）。
図・チャートの理解は不可（本命は NeMo Retriever Library: docs/08 §1・§6A）。

pip install pypdf python-pptx python-docx
使い方: python convert_docs.py --in ./docs_in --out /data/raw --license own
出力: 1要素=1レコード {"text": ..., "meta": {source, source_format, locator, modality, license, provenance}}
"""
from __future__ import annotations
import argparse, json, pathlib, sys, time

def _rec(text, src, fmt, locator, modality, license_, method):
    return {"text": text.strip(),
            "meta": {"source": src, "source_format": fmt, "locator": locator,
                     "modality": modality, "license": license_,
                     "provenance": {"extraction": method,
                                    "created_at": time.strftime("%Y-%m-%dT%H:%M:%S")}}}

def from_pdf(path, license_):
    try:
        from pypdf import PdfReader
    except ImportError:
        sys.exit("pypdf 未導入: pip install pypdf")
    out = []
    for i, page in enumerate(PdfReader(str(path)).pages, 1):
        t = page.extract_text() or ""
        if t.strip():
            out.append(_rec(t, path.name, "pdf", f"p.{i}", "text", license_, "pypdf"))
        else:
            # ネイティブテキストなし=スキャンの可能性 → OCRが必要（本書ではスキップして記録）
            out.append(_rec(f"[SCANNED PAGE p.{i} — OCR required (NeMo Retriever pdfium_hybrid/ocr)]",
                            path.name, "pdf", f"p.{i}", "ocr", license_, "skipped"))
    return out

def _table_to_md(tbl):
    rows = [[(c.text or "").replace("\n", " ").strip() for c in r.cells] for r in tbl.rows]
    if not rows: return ""
    md = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    md += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(md)

def from_pptx(path, license_):
    try:
        from pptx import Presentation
    except ImportError:
        sys.exit("python-pptx 未導入: pip install python-pptx")
    out = []
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides, 1):
        texts, n_pics = [], 0
        for sh in slide.shapes:
            if sh.has_table:
                md = _table_to_md(sh.table)
                if md: out.append(_rec(md, path.name, "pptx", f"slide {i}", "table", license_, "python-pptx"))
            elif getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text)
            elif sh.shape_type == 13:  # PICTURE
                n_pics += 1
        if texts:
            out.append(_rec("\n".join(texts), path.name, "pptx", f"slide {i}", "text", license_, "python-pptx"))
        if n_pics:
            out.append(_rec(f"[{n_pics} FIGURE(S) on slide {i} — VLM caption required (NeMo Retriever .caption())]",
                            path.name, "pptx", f"slide {i}", "figure-caption", license_, "skipped"))
        # 発表者ノート
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame.text
            if nt.strip():
                out.append(_rec(nt, path.name, "pptx", f"slide {i} notes", "speaker-notes", license_, "python-pptx"))
    return out

def from_docx(path, license_):
    try:
        import docx
    except ImportError:
        sys.exit("python-docx 未導入: pip install python-docx")
    d = docx.Document(str(path))
    body = "\n".join(p.text for p in d.paragraphs if p.text.strip())
    out = [_rec(body, path.name, "docx", "body", "text", license_, "python-docx")] if body.strip() else []
    for j, tbl in enumerate(d.tables, 1):
        rows = [[(c.text or "").strip() for c in r.cells] for r in tbl.rows]
        if rows:
            md = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
            md += ["| " + " | ".join(r) + " |" for r in rows[1:]]
            out.append(_rec("\n".join(md), path.name, "docx", f"table {j}", "table", license_, "python-docx"))
    return out

HANDLERS = {".pdf": from_pdf, ".pptx": from_pptx, ".docx": from_docx}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="in_dir", required=True)
    ap.add_argument("--out", default="/data/raw")
    ap.add_argument("--license", dest="license_", default="own")
    a = ap.parse_args()
    out_p = pathlib.Path(a.out); out_p.mkdir(parents=True, exist_ok=True)
    n = skipped = 0
    with (out_p / "docs_extracted.jsonl").open("w", encoding="utf-8") as w:
        for f in sorted(pathlib.Path(a.in_dir).rglob("*")):
            h = HANDLERS.get(f.suffix.lower())
            if not h: continue
            for rec in h(f, a.license_):
                if rec["meta"]["provenance"]["extraction"] == "skipped": skipped += 1
                w.write(json.dumps(rec, ensure_ascii=False) + "\n"); n += 1
    print(f"wrote {n} records -> {out_p/'docs_extracted.jsonl'} (skipped figures/scans: {skipped})")
    if skipped:
        print("⚠ 図/スキャンは未処理。完全な抽出は NeMo Retriever Library を使用（docs/08 §1）")

if __name__ == "__main__":
    main()
