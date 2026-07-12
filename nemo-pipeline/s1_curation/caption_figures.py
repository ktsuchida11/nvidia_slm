"""図・スキャンの処理（フロンティアAPI版）: PDF/PPTXの埋め込み画像を Claude(vision) で
キャプション生成/OCRし、S1入力(jsonl)として出力する。NeMo Retriever(ローカルGPU)の代替ルート。

- 図(figure)      → 日本語キャプション（数値・ラベルは画像から読めるものだけ。推測禁止）
- スキャンページ  → 文字起こし(OCR)
- 幻覚対策: 「読み取れない数値・固有名詞は書かない」を指示し、provenanceにモデル名を記録。
  生成物は"教師出力"扱い → 学習投入前に docs/08 §5 の忠実性ゲート/人手確認を通すこと。

使い方:
  export ANTHROPIC_API_KEY=sk-ant-...
  python caption_figures.py --in ./docs_in --out /data/raw [--model claude-sonnet-4-6] [--max-images 50]
  （--dry-run で API 無しの配管検証）
"""
from __future__ import annotations
import argparse, base64, hashlib, json, pathlib, sys, time

FIG_PROMPT = ("この画像は業務資料中の図表です。内容を日本語で簡潔かつ正確に説明してください。"
              "軸ラベル・凡例・数値は画像から明確に読み取れるものだけを書き、"
              "読み取れない数値・固有名詞は書かないでください。推測や補完は禁止です。")
OCR_PROMPT = ("この画像は文書ページのスキャンです。写っている文字をレイアウト順に日本語で書き起こしてください。"
              "判読できない箇所は『[判読不能]』と記し、推測で補わないでください。")

def sniff_media_type(b: bytes) -> str:
    if b[:8] == b"\x89PNG\r\n\x1a\n": return "image/png"
    if b[:2] == b"\xff\xd8": return "image/jpeg"
    if b[:4] in (b"II*\x00", b"MM\x00*"): return "image/tiff"
    return "image/png"

def call_vlm(model: str, prompt: str, img: bytes) -> str:
    import anthropic  # pip install anthropic
    client = anthropic.Anthropic()
    for attempt in range(3):
        try:
            msg = client.messages.create(
                model=model, max_tokens=700,
                messages=[{"role": "user", "content": [
                    {"type": "image", "source": {"type": "base64",
                        "media_type": sniff_media_type(img),
                        "data": base64.standard_b64encode(img).decode()}},
                    {"type": "text", "text": prompt}]}])
            return msg.content[0].text.strip()
        except Exception as e:
            wait = 2 ** attempt
            print(f"[caption] API error ({e}) retry in {wait}s", file=sys.stderr); time.sleep(wait)
    raise RuntimeError("VLM API failed")

def _rec(text, src, fmt, locator, modality, license_, method):
    return {"text": text.strip(),
            "meta": {"source": src, "source_format": fmt, "locator": locator,
                     "modality": modality, "license": license_,
                     "provenance": {"extraction": method, "method": "teacher-distill",
                                    "created_at": time.strftime("%Y-%m-%dT%H:%M:%S")}}}

def iter_pdf_images(path):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    for i, page in enumerate(reader.pages, 1):
        has_text = bool((page.extract_text() or "").strip())
        for im in page.images:
            yield i, has_text, im.data

def iter_pptx_images(path):
    from pptx import Presentation
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text if slide.shapes.title else ""
        for sh in slide.shapes:
            if sh.shape_type == 13:  # PICTURE
                yield i, title, sh.image.blob

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="in_dir", required=True)
    ap.add_argument("--out", default="/data/raw")
    ap.add_argument("--license", dest="license_", default="own")
    ap.add_argument("--model", default="claude-sonnet-4-6")
    ap.add_argument("--max-images", type=int, default=50, help="コスト暴走防止の上限")
    ap.add_argument("--dry-run", action="store_true", help="API無しの配管検証")
    a = ap.parse_args()
    out_p = pathlib.Path(a.out); out_p.mkdir(parents=True, exist_ok=True)
    method = f"api-vlm:{a.model}" + ("(dry)" if a.dry_run else "")
    cache, n_img, recs = {}, 0, []   # 同一画像はAPI 1回（cache）・レコードは出現箇所ごと

    def caption(img: bytes, prompt: str, placeholder: str) -> str | None:
        nonlocal n_img
        h = hashlib.sha1(img).hexdigest()
        if h in cache:
            return cache[h]                      # 課金せず再利用
        if n_img >= a.max_images:
            print(f"⚠ max-images({a.max_images})到達。以降の新規画像はスキップ")
            return None
        n_img += 1
        cache[h] = placeholder if a.dry_run else call_vlm(a.model, prompt, img)
        return cache[h]

    for f in sorted(pathlib.Path(a.in_dir).rglob("*")):
        try:
            if f.suffix.lower() == ".pdf":
                for pageno, has_text, img in iter_pdf_images(f):
                    modality = "figure-caption" if has_text else "ocr"
                    prompt = FIG_PROMPT if has_text else OCR_PROMPT
                    text = caption(img, prompt, f"[DRY caption p.{pageno}]")
                    if text:
                        recs.append(_rec(text, f.name, "pdf", f"p.{pageno}", modality, a.license_, method))
            elif f.suffix.lower() == ".pptx":
                for slideno, title, img in iter_pptx_images(f):
                    prompt = FIG_PROMPT + (f"（スライドタイトル: {title}）" if title else "")
                    text = caption(img, prompt, f"[DRY caption slide {slideno}]")
                    if text:
                        recs.append(_rec(text, f.name, "pptx", f"slide {slideno}", "figure-caption", a.license_, method))
        except Exception as e:
            print(f"[caption] {f.name}: {e}", file=sys.stderr)

    with (out_p / "figures_captioned.jsonl").open("w", encoding="utf-8") as w:
        for r in recs:
            w.write(json.dumps(r, ensure_ascii=False) + "\n")
    est = n_img * 0.01
    print(f"wrote {len(recs)} records ({n_img} images) -> {out_p/'figures_captioned.jsonl'}")
    print(f"概算コスト: ~${est:.2f}（画像1枚≈$0.01目安・モデル/解像度で変動）" if not a.dry_run else "(dry-run: API未使用)")
    print("★ 生成キャプションは教師出力＝幻覚し得る。学習投入前に忠実性ゲート/サンプリング人手確認(docs/08 §5)")

if __name__ == "__main__":
    main()
